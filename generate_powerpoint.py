import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette Tokens
    NAVY_DARK = RGBColor(15, 23, 42)      # #0F172A - Slide 1 background / dark headers
    SLATE_LIGHT = RGBColor(248, 250, 252) # #F8FAFC - Main slide background
    CARD_BG = RGBColor(255, 255, 255)     # #FFFFFF - Card background
    CARD_BORDER = RGBColor(226, 232, 240) # #E2E8F0 - Subtle hairline border
    CARD_BORDER_DARK = RGBColor(203, 213, 225) # #CBD5E1
    
    PRIMARY_GREEN = RGBColor(16, 185, 129) # #10B981 - AgNER Green
    EMERALD_DARK = RGBColor(5, 150, 105)  # #059669 - Accent dark green
    FOREST_DEEP = RGBColor(6, 78, 59)     # #064E3B - Deep Forest
    
    TEXT_MAIN = RGBColor(30, 41, 59)      # #1E293B - High contrast text
    TEXT_MUTED = RGBColor(100, 116, 139)  # #64748B - Secondary text
    TEXT_LIGHT = RGBColor(241, 245, 249)  # #F1F5F9 - Light text
    
    ACCENT_BLUE = RGBColor(37, 99, 235)   # #2563EB - AI/Model Blue
    ACCENT_CORAL = RGBColor(225, 29, 72)  # #E11D48 - Human / Manual Annotation Red
    ACCENT_AMBER = RGBColor(217, 119, 6)  # #D97706 - Warning / In Progress
    ACCENT_PURPLE = RGBColor(124, 58, 237)# #7C3AED - Quantities / Metadata
    
    FONT_HEADING = "Segoe UI"
    FONT_BODY = "Segoe UI"
    FONT_CODE = "Consolas"

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_top_bar(slide, category, title, subtitle):
        # Category Eyebrow
        eyebrow_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf_e = eyebrow_box.text_frame
        tf_e.word_wrap = True
        tf_e.margin_left = tf_e.margin_top = tf_e.margin_right = tf_e.margin_bottom = 0
        p_e = tf_e.paragraphs[0]
        p_e.text = category.upper()
        p_e.font.size = Pt(10)
        p_e.font.bold = True
        p_e.font.name = FONT_HEADING
        p_e.font.color.rgb = EMERALD_DARK

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.55))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.name = FONT_HEADING
        p_t.font.color.rgb = TEXT_MAIN

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.24), Inches(11.7), Inches(0.35))
        tf_s = sub_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle
        p_s.font.size = Pt(11)
        p_s.font.name = FONT_BODY
        p_s.font.color.rgb = TEXT_MUTED

    def add_footer(slide, current_slide, total_slides=14):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.3))
        tf = footer_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"AgNER Suite • Agricultural Named Entity Recognition & Human-in-the-Loop Workbench    |    Slide {current_slide} of {total_slides}"
        p.font.size = Pt(9)
        p.font.name = FONT_BODY
        p.font.color.rgb = RGBColor(148, 163, 184)

    def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: HERO TITLE SLIDE (Dark Navy)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, NAVY_DARK)

    # Accent glow banner shape
    glow = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.3), Inches(5.8))
    glow.fill.solid()
    glow.fill.fore_color.rgb = PRIMARY_GREEN
    glow.line.fill.background()

    # Title box
    tbox = s1.shapes.add_textbox(Inches(1.4), Inches(1.1), Inches(10.8), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "RESEARCH & ENGINEERING PRESENTATION"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.name = FONT_HEADING
    p0.font.color.rgb = PRIMARY_GREEN

    p1 = tf1.add_paragraph()
    p1.space_before = Pt(12)
    p1.text = "AgNER: Agricultural Named Entity Recognition & Human-in-the-Loop Annotation Suite"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.name = FONT_HEADING
    p1.font.color.rgb = RGBColor(255, 255, 255)

    p2 = tf1.add_paragraph()
    p2.space_before = Pt(14)
    p2.text = "A Complete End-to-End System Architecture: Fine-Tuned RoBERTa (AgNER-BERTa), On-Demand Passage Workbench, 37 Agricultural Taxonomies, and Character-Exact Dataset Engineering."
    p2.font.size = Pt(14)
    p2.font.name = FONT_BODY
    p2.font.color.rgb = RGBColor(203, 213, 225)

    # Key Tech Badges Box
    badges_box = s1.shapes.add_textbox(Inches(1.4), Inches(4.6), Inches(10.8), Inches(1.8))
    tf_b = badges_box.text_frame
    tf_b.word_wrap = True

    pb = tf_b.paragraphs[0]
    pb.text = "TECH STACK & CORE SPECIFICATIONS"
    pb.font.size = Pt(10)
    pb.font.bold = True
    pb.font.color.rgb = PRIMARY_GREEN

    specs = [
        "• Model: RoBERTa-base (512 tokens, 73 logits BIO head, ESWA 2025 Paper Architecture)",
        "• Frontend: React 18, Vite, Tailwind CSS, DOM Range Character-Exact Selection Engine",
        "• Backend & API: Node.js, Express, Mongoose, Sliding Window Prediction, Non-destructive Granular State",
        "• AI Service: FastAPI, PyTorch (CPU Thread-Throttled, micro-batched), Hugging Face Transformers",
        "• Datasets & Storage: Headered JSONL, MongoDB Atlas, CoNLL BIO, Label Studio Tasks, Single-File Enriched Export"
    ]
    for spec in specs:
        p = tf_b.add_paragraph()
        p.space_before = Pt(3)
        p.text = spec
        p.font.size = Pt(10)
        p.font.name = FONT_BODY
        p.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # SLIDE 2: THE PROBLEM STATEMENT & DOMAIN COMPLEXITY
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, SLATE_LIGHT)
    add_top_bar(s2, "Domain Analysis & Motivation", "The Challenge: Extracting Knowledge from Agricultural Advisories", 
                "Why general LLMs and vanilla NLP models fail on domain-specific farmer bulletins and agricultural research documents.")
    add_footer(s2, 2)

    # 3 Cards: Problem 1, Problem 2, Problem 3
    c1 = create_card(s2, 0.8, 1.7, 3.65, 4.9)
    tfc1 = c1.text_frame
    tfc1.word_wrap = True
    tfc1.margin_left = tfc1.margin_top = tfc1.margin_right = tfc1.margin_bottom = Inches(0.25)
    p = tfc1.paragraphs[0]
    p.text = "1. Dense Agronomic Jargon"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = FOREST_DEEP
    points1 = [
        "Mixed Taxonomy: Advisories combine vernacular varieties (PBW 752, Pusa Basmati), chemical actives, local weed names, and meteorological conditions.",
        "Overlapping Categories: A term like 'DAP' or 'Urea' is a Chemical, Fertilizer, and Nutrient depending on context.",
        "Sentence Discontinuity: Traditional sentence splitting breaks advisory contexts where conditions depend on preceding sentences."
    ]
    for pt in points1:
        p = tfc1.add_paragraph()
        p.space_before = Pt(10); p.text = "• " + pt
        p.font.size = Pt(10.5); p.font.color.rgb = TEXT_MAIN

    c2 = create_card(s2, 4.84, 1.7, 3.65, 4.9)
    tfc2 = c2.text_frame
    tfc2.word_wrap = True
    tfc2.margin_left = tfc2.margin_top = tfc2.margin_right = tfc2.margin_bottom = Inches(0.25)
    p = tfc2.paragraphs[0]
    p.text = "2. LLM & Pretrained Limitations"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_AMBER
    points2 = [
        "Hallucinated Boundaries: Generative LLMs (ChatGPT, Gemini) do not emit verified character-level byte/string offsets required for downstream training.",
        "Latency & Financial Cost: Running multi-billion parameter LLMs across hundreds of thousands of state advisory bulletins is computationally prohibitive.",
        "No Ground Truth Verification: Model guesses cannot be treated as gold ground truth without human verification flags."
    ]
    for pt in points2:
        p = tfc2.add_paragraph()
        p.space_before = Pt(10); p.text = "• " + pt
        p.font.size = Pt(10.5); p.font.color.rgb = TEXT_MAIN

    c3 = create_card(s2, 8.88, 1.7, 3.65, 4.9)
    tfc3 = c3.text_frame
    tfc3.word_wrap = True
    tfc3.margin_left = tfc3.margin_top = tfc3.margin_right = tfc3.margin_bottom = Inches(0.25)
    p = tfc3.paragraphs[0]
    p.text = "3. The AgNER Solution"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = PRIMARY_GREEN
    points3 = [
        "Domain-Specialized Encoder: AgNER-BERTa fine-tuned directly on agricultural corpora achieving superior token classification.",
        "Human-in-the-Loop Workbench: On-demand sliding window prediction with zero initial upload wait time.",
        "Character-Exact Engineering: Guaranteed offset fidelity (DOM Range offsets), negative example preservation, and audit-ready data provenance."
    ]
    for pt in points3:
        p = tfc3.add_paragraph()
        p.space_before = Pt(10); p.text = "• " + pt
        p.font.size = Pt(10.5); p.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 3: HIGH-LEVEL SYSTEM ARCHITECTURE
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, SLATE_LIGHT)
    add_top_bar(s3, "Architecture & Engineering", "Tripartite System Architecture: Decoupled Full-Stack Pipeline", 
                "Clean separation of concerns between reactive user interface, stateful REST orchestrator, and high-performance ML inference.")
    add_footer(s3, 3)

    # 4 Vertical Columns representing the tiers
    tiers = [
        ("Client Tier\n(Annotation Workbench)", "React 18 + Vite + Tailwind", [
            "Passage reader with numbered gutter",
            "Left span management deck",
            "Right 37-entity taxonomy rail",
            "DOM Range character offset engine",
            "Optimistic UI & background FIFO queue",
            "One-click dataset exporter"
        ], PRIMARY_GREEN),
        ("API & State Tier\n(Backend Orchestrator)", "Node.js + Express + Mongoose", [
            "Document ingestion & header parsing",
            "Sliding-window prediction manager",
            "Granular line-level persistence (PUT)",
            "Immutable model resets (POST)",
            "DatasetBuilder (BIO, HF, CoNLL)",
            "Round-trip import validation"
        ], ACCENT_BLUE),
        ("Inference Tier\n(AgNER AI Engine)", "FastAPI + PyTorch + HF", [
            "RoBERTa-base token classification",
            "512 token context windowing",
            "200-word chunking & subword mapping",
            "Softmax token-level confidence averaging",
            "CPU multi-thread throttling (2 vCPU)",
            "Micro-batching (5-passage chunks)"
        ], ACCENT_PURPLE),
        ("Storage Tier\n(Persistence & Artifacts)", "MongoDB Atlas + JSONL", [
            "Document & Segment schemas",
            "Immutable originalEntities copy",
            "Active entities with provenance tags",
            "Headered JSONL advisory files",
            "Label Studio export XML/JSON",
            "Versioned data manifest (schema 1.1)"
        ], EMERALD_DARK),
    ]

    for i, (title, subtitle, bullets, color) in enumerate(tiers):
        left = 0.8 + i * 2.98
        c = create_card(s3, left, 1.7, 2.8, 4.9)
        tfc = c.text_frame
        tfc.word_wrap = True
        tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = Inches(0.2)
        
        # Pill header
        p = tfc.paragraphs[0]
        p.text = title
        p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = color
        
        p_sub = tfc.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(9.5); p_sub.font.bold = True; p_sub.font.color.rgb = TEXT_MUTED
        
        for b in bullets:
            pb = tfc.add_paragraph()
            pb.space_before = Pt(8)
            pb.text = "• " + b
            pb.font.size = Pt(9.5); pb.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 4: DEEP DIVE: THE AI MODEL (AgNER-BERTa)
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, SLATE_LIGHT)
    add_top_bar(s4, "Machine Learning Engine", "AgNER-BERTa: Transformer Token Classification for Agronomy", 
                "Grounded in published research: fine-tuned encoder models with domain data augmentation outperform ChatGPT in Agri-NER.")
    add_footer(s4, 4)

    # Left card: Model Architecture & Specs
    c_m1 = create_card(s4, 0.8, 1.7, 5.7, 4.9)
    tf_m1 = c_m1.text_frame
    tf_m1.word_wrap = True
    tf_m1.margin_left = tf_m1.margin_top = tf_m1.margin_right = tf_m1.margin_bottom = Inches(0.25)
    
    p = tf_m1.paragraphs[0]
    p.text = "Model Specifications & Token Head"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_BLUE

    m_specs = [
        ("Base Backbone", "RoBERTa-base (12 layers, 768 hidden dimension, 12 attention heads, ~125M parameters)"),
        ("Max Sequence Length", "512 subword tokens per forward pass with sliding 200-word chunking window"),
        ("Classification Head", "73 Linear Output Logits: 'O' (Outside) + 36 Entity Classes × 2 (B-Begin, I-Inside)"),
        ("Token-to-Word Alignment", "Encodes split words with add_prefix_space=True; assigns first subtoken prediction to word"),
        ("Confidence Metric", "Softmax probability over predicted class, averaged across constituent entity tokens"),
        ("Scientific Foundation", "Published in Expert Systems with Applications (ESWA 2025): 'Fine-tuned encoder models with data augmentation beat ChatGPT in agricultural named entity recognition'"),
    ]
    for label, val in m_specs:
        p = tf_m1.add_paragraph()
        p.space_before = Pt(8)
        p.text = f"{label}: "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MAIN
        # append value
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right card: Inference Optimization & Boundary Alignment
    c_m2 = create_card(s4, 6.8, 1.7, 5.7, 4.9)
    tf_m2 = c_m2.text_frame
    tf_m2.word_wrap = True
    tf_m2.margin_left = tf_m2.margin_top = tf_m2.margin_right = tf_m2.margin_bottom = Inches(0.25)
    
    p = tf_m2.paragraphs[0]
    p.text = "Inference Pipeline & Boundary Sanitation"
    p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = PRIMARY_GREEN

    m_pipe = [
        ("Chunking Algorithm", "Regex word splitting with finditer(r'\\S+'). Chunks capped at CHUNK_WORDS (200 words) to guarantee < 512 subwords without truncation loss."),
        ("No-Grad Evaluation", "Executed under @torch.inference_mode() with low_cpu_mem_usage=True, avoiding graph allocation."),
        ("Punctuation Trimming", "Heuristic post-processing automatically contracts trailing sentence delimiters (.,;:!?)\"'). Words like 'yellow rust.' cleanly isolate 'yellow rust'."),
        ("Multi-Thread Throttling", "torch.set_num_threads(2) & interop=1 to prevent thread contention and vCPU saturation in shared cloud container environments."),
        ("Provenance Injection", "Output entities tagged with tag: 'NER_LABELLED', labelled_by: 'NER Model', source: 'model', status: 'accepted' with calibrated float confidence.")
    ]
    for label, val in m_pipe:
        p = tf_m2.add_paragraph()
        p.space_before = Pt(8)
        p.text = f"{label}: "
        p.font.bold = True
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_MAIN
        run = p.add_run()
        run.text = val
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 5: 37 CANONICAL AGRICULTURAL TAXONOMIES
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, SLATE_LIGHT)
    add_top_bar(s5, "Domain Taxonomy & Ontology", "37 Canonical Agricultural Entities Across 7 Functional Domains", 
                "Standardized schema ensuring comprehensive coverage from agronomy and pathology to weather, policies, and farm equipment.")
    add_footer(s5, 5)

    # Create 7 Domain Cards arranged in a 2-column or 3-column grid
    domains = [
        ("Crops & Agriculture (5)", PRIMARY_GREEN, "Crop, Fruit, Vegetable, Food_Item, Soil", "Wheat, PBW 752, Maize, Basmati, Sandy Loam, Clay Soil"),
        ("Pests & Pathology (4)", ACCENT_CORAL, "Disease, Organism, Agri_Pollution, Agri_Waste", "Yellow Rust, Stem Borer, Aphids, Whitefly, Stubble Burning"),
        ("Treatments & Inputs (3)", ACCENT_BLUE, "Chemical, Treatment, Nutrient", "Urea, DAP, Imidacloprid, Fungicide Spray, Nitrogen, Potash"),
        ("Meteorology & Climate (7)", RGBColor(14, 165, 233), "Weather, Rainfall, Temp, Humidity, Season, Natural_Disaster, Natural_Resource", "Monsoon, Frost, 28°C, 85% RH, Kharif, Drought, Groundwater"),
        ("Quantities & Time (7)", ACCENT_PURPLE, "Quantity, Other_Quantity, Field_Area, Duration, Money, Date_and_Time, Event", "55 kg DAP, 2.5 acres, 15 days, ₹1200, Mid-December, Sowing"),
        ("Entities & Agents (4)", RGBColor(79, 70, 229), "Person, Organization, Location, Policy", "PAU Ludhiana, ICAR, Punjab, Extension Officer, PM-KISAN"),
        ("Methods & Machinery (7)", ACCENT_AMBER, "Agri_Equipment, Agri_Method, Agri_Process, Technology, ML_Model, Citation, Other", "Happy Seeder, Super Seeder, Rotavator, Drip Irrigation, Tillering")
    ]

    # Grid of cards: 4 on top, 3 on bottom
    for i in range(4):
        d_name, d_col, d_classes, d_ex = domains[i]
        c = create_card(s5, 0.8 + i * 2.98, 1.7, 2.8, 2.35)
        tf = c.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.15)
        p = tf.paragraphs[0]; p.text = d_name; p.font.size = Pt(11.5); p.font.bold = True; p.font.color.rgb = d_col
        p2 = tf.add_paragraph(); p2.space_before = Pt(4); p2.text = f"Classes: {d_classes}"; p2.font.size = Pt(9); p2.font.bold = True; p2.font.color.rgb = TEXT_MAIN
        p3 = tf.add_paragraph(); p3.space_before = Pt(4); p3.text = f"Examples: {d_ex}"; p3.font.size = Pt(8.5); p3.font.color.rgb = TEXT_MUTED

    for i in range(3):
        d_name, d_col, d_classes, d_ex = domains[4 + i]
        c = create_card(s5, 0.8 + i * 3.98, 4.25, 3.8, 2.35)
        tf = c.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.15)
        p = tf.paragraphs[0]; p.text = d_name; p.font.size = Pt(11.5); p.font.bold = True; p.font.color.rgb = d_col
        p2 = tf.add_paragraph(); p2.space_before = Pt(4); p2.text = f"Classes: {d_classes}"; p2.font.size = Pt(9); p2.font.bold = True; p2.font.color.rgb = TEXT_MAIN
        p3 = tf.add_paragraph(); p3.space_before = Pt(4); p3.text = f"Examples: {d_ex}"; p3.font.size = Pt(8.5); p3.font.color.rgb = TEXT_MUTED

    # Special callout box inside bottom right card or highlight
    callout = create_card(s5, 8.88, 4.25, 3.65, 2.35, bg_color=RGBColor(254, 242, 242), border_color=ACCENT_CORAL)
    tfc = callout.text_frame; tfc.word_wrap = True; tfc.margin_left = tfc.margin_top = tfc.margin_right = tfc.margin_bottom = Inches(0.15)
    p = tfc.paragraphs[0]; p.text = "Key Architectural Detail: Agri_Equipment"; p.font.size = Pt(11.5); p.font.bold = True; p.font.color.rgb = ACCENT_CORAL
    p2 = tfc.add_paragraph(); p2.space_before = Pt(4)
    p2.text = "• 37th Entity Class added for farm machinery (happy seeder, super seeder, rotavator).\n• Model checkpoint head remains 36 classes (73 logits). Agri_Equipment is an 'annotator-only' class.\n• In exports, always tagged with source: 'human', verified: true, ensuring zero model drift while capturing vital tractor & machinery entities."
    p2.font.size = Pt(8.5); p2.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 6: DOCUMENT INGESTION & PARSING ENGINE
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, SLATE_LIGHT)
    add_top_bar(s6, "Data Ingestion Pipeline", "Advisory Ingestion & Zero-Wait Headered JSONL Parser", 
                "Reduced file upload time from several minutes to 67ms by completely decoupling ingestion from model inference.")
    add_footer(s6, 6)

    c_i1 = create_card(s6, 0.8, 1.7, 5.7, 4.9)
    tf_i1 = c_i1.text_frame; tf_i1.word_wrap = True; tf_i1.margin_left = tf_i1.margin_top = tf_i1.margin_right = tf_i1.margin_bottom = Inches(0.25)
    p = tf_i1.paragraphs[0]; p.text = "The Ingestion Architecture Overhaul"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = EMERALD_DARK

    ingest_pts = [
        ("Old Bottleneck", "Legacy systems ran full model predictions at document upload time. A 323-passage file froze the UI for 3-5 minutes, leading to proxy timeouts (502 Bad Gateway) and wasted GPU/CPU cycles on unviewed passages."),
        ("New Decoupled Pipeline", "Upload only parses text and header metadata into MongoDB, marking status: 'pending'. The 323-passage file now opens in 67 milliseconds."),
        ("First-Class Advisory Metadata", "Extracts passage_id, passage_tag, crops, stages, and subjects. Carries them intact through UI display down to final export files."),
        ("Universal Format Support", "One text field per record is the only strict requirement. Automatically parses nested advisory sections, paged documents, test_cases, and raw JSONL.")
    ]
    for title, desc in ingest_pts:
        p = tf_i1.add_paragraph(); p.space_before = Pt(8)
        p.text = f"{title}: "; p.font.bold = True; p.font.size = Pt(10); p.font.color.rgb = TEXT_MAIN
        run = p.add_run(); run.text = desc; run.font.bold = False; run.font.color.rgb = TEXT_MUTED

    # Right Card: Sample Record JSON & Parsing Flow
    c_i2 = create_card(s6, 6.8, 1.7, 5.7, 4.9)
    tf_i2 = c_i2.text_frame; tf_i2.word_wrap = True; tf_i2.margin_left = tf_i2.margin_top = tf_i2.margin_right = tf_i2.margin_bottom = Inches(0.25)
    p = tf_i2.paragraphs[0]; p.text = "Advisory Record Schema & Normalization"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_BLUE

    schema_text = """// Ingested JSONL Advisory Record:
{
  "passage_id": "3817e6a6596d",
  "passage_tag": "PUN-0001",
  "crops": ["WHEAT"],
  "stages": ["Sowing"],
  "subjects": ["WHEAT (Sowing)"],
  "text": "Farmers are advised to sow PBW 752 with Happy Seeder..."
}

// Normalized Internal Segment Document:
{
  "_id": ObjectId("..."),
  "segmentIndex": 0,
  "passageId": "3817e6a6596d",
  "passageTag": "PUN-0001",
  "crops": ["WHEAT"],
  "stages": ["Sowing"],
  "nerStatus": "pending",        // Predicted on demand
  "sentences": [
    { "lineIndex": 0, "text": "...", "entities": [] }
  ]
}"""
    p_code = tf_i2.add_paragraph(); p_code.space_before = Pt(6)
    p_code.text = schema_text; p_code.font.size = Pt(8.5); p_code.font.name = FONT_CODE; p_code.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 7: ON-DEMAND STREAMING PREDICTION PIPELINE
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, SLATE_LIGHT)
    add_top_bar(s7, "Inference Orchestration", "On-Demand Windowed Inference & Asynchronous Writes", 
                "The model runs in a dynamic sliding window 10 passages ahead of the annotator, with zero network blocking.")
    add_footer(s7, 7)

    # 3 Workflow stages horizontally
    steps = [
        ("1. Sliding Window Prediction", ACCENT_BLUE, [
            "POST /api/documents/:id/predict",
            "Client requests active passage + next 10 passages ({ from: index, count: 10 }).",
            "Server only queries passages with status 'pending' or 'error'.",
            "Passages already marked or verified by human are never overwritten.",
            "Failed inference windows are immediately retryable with 1-click."
        ]),
        ("2. FIFO Background Write Queue", PRIMARY_GREEN, [
            "Optimistic Local Updates",
            "Every span addition, category change, or deletion reflects instantly on UI.",
            "Write requests enter a strictly ordered client-side FIFO queue.",
            "Prevents race conditions in segment-level read-modify-write transactions.",
            "Returns compact status counts instead of echoing 12.8MB document payloads."
        ]),
        ("3. Granular State Preservation", ACCENT_PURPLE, [
            "Immutable Original Copy",
            "Model predictions saved to immutable originalEntities field in MongoDB.",
            "Line-level reset (POST .../reset-line) restores AI predictions for one sentence.",
            "Arrow keys (< / >) auto-commit pending edits before switching passages."
        ])
    ]

    for i, (stitle, scol, sbullets) in enumerate(steps):
        c = create_card(s7, 0.8 + i * 3.98, 1.7, 3.8, 4.9)
        tf = c.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.2)
        p = tf.paragraphs[0]; p.text = stitle; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = scol
        for b in sbullets:
            pb = tf.add_paragraph(); pb.space_before = Pt(10)
            pb.text = "• " + b; pb.font.size = Pt(10); pb.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 8: PASSAGE ANNOTATION WORKBENCH (FRONTEND)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, SLATE_LIGHT)
    add_top_bar(s8, "Frontend UI & UX Architecture", "The Passage Annotation Workbench: 3-Column Ergonomics", 
                "Designed around the full passage rather than isolated sentences, maintaining agronomic context and rapid keyboard workflows.")
    add_footer(s8, 8)

    # 3 Columns representing the UI
    ui_cols = [
        ("Left Panel: Spans Deck", "SpanListPanel.jsx", ACCENT_PURPLE, [
            "Live Spans Roster: Lists every entity in the current passage.",
            "Sentence Index Badges: Displays which sentence line each span belongs to.",
            "One-Click Actions: Approve individual spans, 'Approve All' passage spans, or delete false positives.",
            "Provenance Badges: Clearly distinguishes model-proposed vs human-added entities."
        ]),
        ("Center Panel: Passage Reader", "PassageAnnotationCard.jsx", PRIMARY_GREEN, [
            "Gutter-Numbered Sentences: Displays all sentences in the advisory consecutively (20px serif font).",
            "Dual-Style Underlines: Solid colored underline for model proposals; dashed underline for human tags.",
            "Tag Captions: Rendered below spans with entity type and remove button.",
            "Context Headers: Displays PUN advisory tag, target crops, growth stages, and subjects."
        ]),
        ("Right Rail: 37 Entity Palette", "EntityPaletteSidebar.jsx", ACCENT_BLUE, [
            "Always-Open Rail: Eliminates hidden dropdowns; all 37 classes always accessible.",
            "Dynamic Recently Used: Follows real usage patterns rather than fixed defaults.",
            "Numeric Shortcuts: 1-9 apply key labels immediately without mouse movement.",
            "Live Typeahead Search: Press '/' to trigger instant filter across all categories."
        ])
    ]

    for i, (title, comp, col, bullets) in enumerate(ui_cols):
        c = create_card(s8, 0.8 + i * 3.98, 1.7, 3.8, 4.9)
        tf = c.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.2)
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(14); p.font.bold = True; p.font.color.rgb = col
        p_sub = tf.add_paragraph(); p_sub.text = f"Component: {comp}"; p_sub.font.size = Pt(9); p_sub.font.bold = True; p_sub.font.color.rgb = TEXT_MUTED
        for b in bullets:
            pb = tf.add_paragraph(); pb.space_before = Pt(10)
            pb.text = "• " + b; pb.font.size = Pt(10); pb.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 9: CHARACTER-EXACT OFFSET ENGINE
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, SLATE_LIGHT)
    add_top_bar(s9, "Algorithms & Precision", "The Exact Character Offset Engine (`selectionOffsets.js`)", 
                "Eliminating string-search offset drift and repeated-token collision through DOM Range subtree measurement.")
    add_footer(s9, 9)

    c_o1 = create_card(s9, 0.8, 1.7, 5.7, 4.9)
    tf_o1 = c_o1.text_frame; tf_o1.word_wrap = True; tf_o1.margin_left = tf_o1.margin_top = tf_o1.margin_right = tf_o1.margin_bottom = Inches(0.25)
    p = tf_o1.paragraphs[0]; p.text = "The Legacy Fallback Flaw vs New Engine"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_CORAL

    flaws = [
        ("The Legacy Problem", "Old tools estimated offsets by measuring rendered text before selection and hunting within 50 characters, falling back to the first occurrence in the sentence."),
        ("Caption Pollution", "The small tag captions ('Crop ×') rendered inside the container. An 8-span sentence was 81 characters off. Fallbacks repeatedly tagged the wrong occurrence of repeated words (e.g. multiple 'wheat' tokens)."),
        ("The New Algorithm", "Clones the DOM Range preceding the selection, recursively strips all elements marked with data-caption, and measures the exact character length."),
        ("select-none Protection", "Captions marked select-none so user selections never drag across tag labels."),
        ("Guaranteed Alignment", "Span text is sliced directly from sentence.text at the measured offsets. Text and offsets cannot disagree.")
    ]
    for title, desc in flaws:
        p = tf_o1.add_paragraph(); p.space_before = Pt(8)
        p.text = f"{title}: "; p.font.bold = True; p.font.size = Pt(10); p.font.color.rgb = TEXT_MAIN
        run = p.add_run(); run.text = desc; run.font.bold = False; run.font.color.rgb = TEXT_MUTED

    # Right Card: Code Algorithm snippet
    c_o2 = create_card(s9, 6.8, 1.7, 5.7, 4.9)
    tf_o2 = c_o2.text_frame; tf_o2.word_wrap = True; tf_o2.margin_left = tf_o2.margin_top = tf_o2.margin_right = tf_o2.margin_bottom = Inches(0.25)
    p = tf_o2.paragraphs[0]; p.text = "DOM Range Offset Measurement Algorithm"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = PRIMARY_GREEN

    algo_code = """export function getSelectionOffsets(rootNode, rawSentenceText) {
  const selection = window.getSelection();
  if (!selection.rangeCount) return null;
  const range = selection.getRangeAt(0);

  // 1. Measure start offset: clone range from root start to selection start
  const preRange = range.cloneRange();
  preRange.selectNodeContents(rootNode);
  preRange.setEnd(range.startContainer, range.startOffset);
  
  // 2. Clone fragment & remove all caption subtrees ([data-caption])
  const fragment = preRange.cloneContents();
  fragment.querySelectorAll('[data-caption]').forEach(el => el.remove());
  const start = fragment.textContent.length;
  
  // 3. Measure selected text fragment length (excluding captions)
  const selFragment = range.cloneContents();
  selFragment.querySelectorAll('[data-caption]').forEach(el => el.remove());
  const length = selFragment.textContent.length;
  const end = start + length;

  // 4. Extract exact substring from authoritative text
  return { start, end, text: rawSentenceText.slice(start, end) };
}"""
    p_code = tf_o2.add_paragraph(); p_code.space_before = Pt(6)
    p_code.text = algo_code; p_code.font.size = Pt(8); p_code.font.name = FONT_CODE; p_code.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 10: GRANULAR PERSISTENCE & STATE MANAGEMENT
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, SLATE_LIGHT)
    add_top_bar(s10, "Persistence & Database Architecture", "Granular State Persistence & Non-Destructive Storage", 
                "Ensuring high concurrency, non-destructive line edits, and sub-millisecond local response times.")
    add_footer(s10, 10)

    # 2 Big Cards: MongoDB Architecture and Controller Mechanics
    c_p1 = create_card(s10, 0.8, 1.7, 5.7, 4.9)
    tf_p1 = c_p1.text_frame; tf_p1.word_wrap = True; tf_p1.margin_left = tf_p1.margin_top = tf_p1.margin_right = tf_p1.margin_bottom = Inches(0.25)
    p = tf_p1.paragraphs[0]; p.text = "MongoDB Schema Architecture"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_BLUE

    schema_pts = [
        ("Two-Tier Model", "Document model stores top-level metadata, statistics, and segment references. Segment model stores sentence arrays and entities."),
        ("Dual Entity Arrays", "Each sentence maintains originalEntities (immutable model cache) and entities (active mutable annotations with provenance)."),
        ("Provenance Tracking", "Every entity records start, end, label, text, source ('model' | 'human'), verified (boolean), status ('accepted' | 'approved' | 'edited'), and confidence."),
        ("Compact Write Payloads", "Instead of returning the entire document tree (12.8MB in large files), API endpoints return concise status codes and updated entity counts.")
    ]
    for title, desc in schema_pts:
        p = tf_p1.add_paragraph(); p.space_before = Pt(8)
        p.text = f"{title}: "; p.font.bold = True; p.font.size = Pt(10); p.font.color.rgb = TEXT_MAIN
        run = p.add_run(); run.text = desc; run.font.bold = False; run.font.color.rgb = TEXT_MUTED

    c_p2 = create_card(s10, 6.8, 1.7, 5.7, 4.9)
    tf_p2 = c_p2.text_frame; tf_p2.word_wrap = True; tf_p2.margin_left = tf_p2.margin_top = tf_p2.margin_right = tf_p2.margin_bottom = Inches(0.25)
    p = tf_p2.paragraphs[0]; p.text = "Controller Execution Logic"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = EMERALD_DARK

    ctrl_pts = [
        ("PUT .../line-entities", "Commits human additions, removals, and relabels for a single sentence line index without locking or modifying adjacent sentences."),
        ("POST .../reset-line", "Reverts a single line back to original AI predictions by copying originalEntities to entities, restoring accidental deletions."),
        ("Unsaved Changes Protection", "Navigating with keyboard arrows (Left/Right) automatically flushes pending edits to MongoDB before switching passage contexts."),
        ("Batch Window Updates", "predictController executes windowed bulk updates directly on Segment models using Mongoose bulkWrite() for speed.")
    ]
    for title, desc in ctrl_pts:
        p = tf_p2.add_paragraph(); p.space_before = Pt(8)
        p.text = f"{title}: "; p.font.bold = True; p.font.color.rgb = TEXT_MAIN
        run = p.add_run(); run.text = desc; run.font.bold = False; run.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 11: DATASET BUILDER & EXPORT STATION
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, SLATE_LIGHT)
    add_top_bar(s11, "Export & Machine Learning Formats", "DatasetBuilder: Production Multi-Format Export Station", 
                "Exports authoritative single-file datasets that serve both fine-tuning and honest gold-standard evaluation.")
    add_footer(s11, 11)

    # 4 Cards across
    export_formats = [
        ("Enriched Single-File Export", "GET /api/documents/:id/dataset", PRIMARY_GREEN, [
            "Returns uploaded JSONL with original fields in identical order.",
            "Appends entities array and annotation summary block.",
            "Negative Examples Preserved: Passages with no entities remain (essential for training).",
            "Round-Trip Re-import: Exported file can be re-uploaded to continue work seamlessly."
        ]),
        ("Hugging Face Datasets", "?format=hf", ACCENT_BLUE, [
            "Tokens array with tokenized subwords/words.",
            "ner_tags (string BIO tags: 'B-Crop', 'I-Crop').",
            "ner_ids (integer IDs dynamically mapped to bio_label_to_id from manifest).",
            "Ready for Hugging Face Trainer & evaluate."
        ]),
        ("CoNLL BIO & SpaCy", "?format=conll | spacy", ACCENT_PURPLE, [
            "Tab-separated word and BIO tag formatting.",
            "Conflict Resolution: Overlaps resolved by human-verified first, then length, then confidence.",
            "Token-safe hyphens: 'mid-December' split to keep 'December' span valid in BIO."
        ]),
        ("Label Studio Official Tasks", "GET .../label-studio", ACCENT_AMBER, [
            "Official Label Studio JSON schema.",
            "Pre-annotations section with AgNER-BERTa scores.",
            "Annotations section with human verified tags.",
            "One-click import into enterprise Label Studio setups."
        ])
    ]

    for i, (title, endpoint, col, bullets) in enumerate(export_formats):
        c = create_card(s11, 0.8 + i * 2.98, 1.7, 2.8, 4.9)
        tf = c.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.2)
        p = tf.paragraphs[0]; p.text = title; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = col
        p_sub = tf.add_paragraph(); p_sub.text = endpoint; p_sub.font.size = Pt(8.5); p_sub.font.name = FONT_CODE; p_sub.font.bold = True; p_sub.font.color.rgb = TEXT_MUTED
        for b in bullets:
            pb = tf.add_paragraph(); pb.space_before = Pt(8)
            pb.text = "• " + b; pb.font.size = Pt(9.5); pb.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 12: CONTAINERIZATION & CLOUD DEPLOYMENT
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, SLATE_LIGHT)
    add_top_bar(s12, "DevOps & Cloud Engineering", "Unified Production Containerization & Cloud Deployment", 
                "Single multi-stage Docker container orchestrating Node.js, FastAPI, PyTorch, and hardened networking.")
    add_footer(s12, 12)

    c_d1 = create_card(s12, 0.8, 1.7, 5.7, 4.9)
    tf_d1 = c_d1.text_frame; tf_d1.word_wrap = True; tf_d1.margin_left = tf_d1.margin_top = tf_d1.margin_right = tf_d1.margin_bottom = Inches(0.25)
    p = tf_d1.paragraphs[0]; p.text = "Container Orchestration Architecture"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = EMERALD_DARK

    devops_pts = [
        ("Unified Single Container", "Bundles Node.js runtime and Python 3.10+ in one production image (NER-App/Dockerfile) managed by start.sh startup script."),
        ("Micro-Batching", "Segment prediction requests split into chunks of 5 segments, keeping latency well under reverse proxy 60-second timeouts (eliminating 502 Bad Gateway)."),
        ("Hardened Dynamic CORS", "Configured with dynamic origin reflection (origin: true), preflight catch-all (app.options('*')), and error middleware header persistence."),
        ("Zero-Dependency Fast Boot", "Model safetensors loaded via AutoModelForTokenClassification.from_pretrained() with local cache fallback.")
    ]
    for title, desc in devops_pts:
        p = tf_d1.add_paragraph(); p.space_before = Pt(8)
        p.text = f"{title}: "; p.font.bold = True; p.font.size = Pt(10); p.font.color.rgb = TEXT_MAIN
        run = p.add_run(); run.text = desc; run.font.bold = False; run.font.color.rgb = TEXT_MUTED

    # Right Card: start.sh and Dockerfile structure
    c_d2 = create_card(s12, 6.8, 1.7, 5.7, 4.9)
    tf_d2 = c_d2.text_frame; tf_d2.word_wrap = True; tf_d2.margin_left = tf_d2.margin_top = tf_d2.margin_right = tf_d2.margin_bottom = Inches(0.25)
    p = tf_d2.paragraphs[0]; p.text = "start.sh Container Startup Orchestration"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_BLUE

    docker_code = """#!/usr/bin/env bash
set -e

echo "Starting AgNER Production Container Services..."

# 1. Start Python FastAPI AI Inference Service
export TORCH_THREADS=2
export CHUNK_WORDS=200
cd /app/NER-App/python-ai-service
uvicorn app:app --host 127.0.0.1 --port 8000 &
AI_PID=$!

# 2. Wait for AI Service /health check
until curl -s http://127.0.0.1:8000/health > /dev/null; do
  echo "Waiting for AgNER-BERTa checkpoint to load..."
  sleep 1
done
echo "AgNER-BERTa AI Engine Online (36 Model Classes)"

# 3. Start Node.js Express Backend
cd /app/NER-App/backend
node src/server.js"""
    p_code = tf_d2.add_paragraph(); p_code.space_before = Pt(6)
    p_code.text = docker_code; p_code.font.size = Pt(8.5); p_code.font.name = FONT_CODE; p_code.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 13: END-TO-END WORKFLOW & USER JOURNEY
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, SLATE_LIGHT)
    add_top_bar(s13, "Operational Workflow", "End-to-End Operational Lifecycle: From Raw Advisory to Gold Dataset", 
                "The complete step-by-step lifecycle of an agricultural advisory document in the AgNER platform.")
    add_footer(s13, 13)

    # 5 Process Steps horizontally
    flow_steps = [
        ("Step 1: Upload", "Raw Advisory Ingestion", [
            "Upload .jsonl or .json advisory file.",
            "Parses text and Punjab PAU metadata in 67ms.",
            "Passages saved to MongoDB as 'pending'."
        ]),
        ("Step 2: Stream", "On-Demand Window", [
            "Workbench opens passage 1 immediately.",
            "AI service predicts passages 1-10 ahead.",
            "AgNER-BERTa proposes 36 classes with confidence."
        ]),
        ("Step 3: Review", "Passage Workbench", [
            "Annotator reviews numbered sentences.",
            "Solid underline: model proposals.",
            "Dashed underline: manual human markup.",
            "Quick hotkeys: 1-9 label, '/' searches."
        ]),
        ("Step 4: Verify", "Granular Auditing", [
            "Hit 'A' to Approve entire passage.",
            "Delete false positives, add missing seeders.",
            "Edits queued and saved via FIFO stream."
        ]),
        ("Step 5: Export", "Dataset Generation", [
            "One-click direct browser file download.",
            "Emits original JSONL with verified entities.",
            "Preserves negative samples & metadata.",
            "Ready for CoNLL / HF Trainer fine-tuning."
        ])
    ]

    for i, (st, sub, bullets) in enumerate(flow_steps):
        c = create_card(s13, 0.8 + i * 2.38, 1.7, 2.25, 4.9)
        tf = c.text_frame; tf.word_wrap = True; tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.15)
        p = tf.paragraphs[0]; p.text = st; p.font.size = Pt(13); p.font.bold = True; p.font.color.rgb = PRIMARY_GREEN
        p_sub = tf.add_paragraph(); p_sub.text = sub; p_sub.font.size = Pt(9); p_sub.font.bold = True; p_sub.font.color.rgb = TEXT_MUTED
        for b in bullets:
            pb = tf.add_paragraph(); pb.space_before = Pt(8)
            pb.text = "• " + b; pb.font.size = Pt(9); pb.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 14: SUMMARY, RESEARCH IMPACT & FUTURE ROADMAP
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, SLATE_LIGHT)
    add_top_bar(s14, "Summary & Future Roadmap", "Summary of Achievements, Research Impact & Next Horizons", 
                "Bridging the gap between cutting-edge NLP research and real-world agricultural field advisories.")
    add_footer(s14, 14)

    # 3 Cards
    c_s1 = create_card(s14, 0.8, 1.7, 3.65, 4.9)
    tf_s1 = c_s1.text_frame; tf_s1.word_wrap = True; tf_s1.margin_left = tf_s1.margin_top = tf_s1.margin_right = tf_s1.margin_bottom = Inches(0.25)
    p = tf_s1.paragraphs[0]; p.text = "Engineering Milestones"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = PRIMARY_GREEN
    s1_pts = [
        "Zero-Wait Ingestion: From minutes to 67ms for 300+ passage files.",
        "Exact Character Offsets: Zero alignment errors on repeated tokens via DOM Range engine.",
        "Non-Destructive Persistence: Fine-grained sentence-level saves and immutable resets.",
        "Human Provenance Tracking: Explicitly recording whether each span was verified by a human or generated by AI."
    ]
    for pt in s1_pts:
        p = tf_s1.add_paragraph(); p.space_before = Pt(10)
        p.text = "• " + pt; p.font.size = Pt(10.5); p.font.color.rgb = TEXT_MAIN

    c_s2 = create_card(s14, 4.84, 1.7, 3.65, 4.9)
    tf_s2 = c_s2.text_frame; tf_s2.word_wrap = True; tf_s2.margin_left = tf_s2.margin_top = tf_s2.margin_right = tf_s2.margin_bottom = Inches(0.25)
    p = tf_s2.paragraphs[0]; p.text = "Scientific Contribution"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_BLUE
    s2_pts = [
        "Beating General LLMs: Confirmed ESWA 2025 findings that fine-tuned encoders with domain augmentation outperform generative LLMs in agricultural NER.",
        "Auditable Gold Datasets: Single-file export distinguishing verified spans from raw AI output.",
        "Negative Sample Preservation: Preventing model over-prediction bias by retaining entity-free sentences."
    ]
    for pt in s2_pts:
        p = tf_s2.add_paragraph(); p.space_before = Pt(10)
        p.text = "• " + pt; p.font.size = Pt(10.5); p.font.color.rgb = TEXT_MAIN

    c_s3 = create_card(s14, 8.88, 1.7, 3.65, 4.9)
    tf_s3 = c_s3.text_frame; tf_s3.word_wrap = True; tf_s3.margin_left = tf_s3.margin_top = tf_s3.margin_right = tf_s3.margin_bottom = Inches(0.25)
    p = tf_s3.paragraphs[0]; p.text = "Future Project Roadmap"; p.font.size = Pt(15); p.font.bold = True; p.font.color.rgb = ACCENT_PURPLE
    s3_pts = [
        "Active Learning Loop: Auto-triggering model fine-tuning checkpoints when verified spans cross volume thresholds.",
        "Relation Extraction (RE): Marking relations between entities (e.g. Chemical 'Urea' -> applied_to -> Crop 'Wheat').",
        "Multilingual Advisory Support: Expanding tokenizer and models to Hindi, Punjabi (Gurmukhi), and Marathi state bulletins."
    ]
    for pt in s3_pts:
        p = tf_s3.add_paragraph(); p.space_before = Pt(10)
        p.text = "• " + pt; p.font.size = Pt(10.5); p.font.color.rgb = TEXT_MAIN

    # Save presentation
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    out_file = sys.argv[1] if len(sys.argv) > 1 else "AgNER_Technical_Presentation.pptx"
    create_deck(out_file)
